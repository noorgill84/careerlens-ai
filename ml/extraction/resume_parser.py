"""
Resume parsing: PDF/DOCX -> structured ResumeProfile.

Pipeline (per spec §7):
  bytes -> raw text -> clean_text -> section_detector -> field extraction
  -> skill_extractor -> ResumeProfile (skills not yet normalized here;
  normalization happens in skill_taxonomy.normalize_skills at the caller,
  keeping this module focused on extraction only).
"""
from __future__ import annotations
import io
import re
from dataclasses import dataclass, field

from ml.preprocessing.text_cleaning import clean_text, extract_contact_fields, split_lines
from ml.extraction.section_detector import detect_sections
from ml.extraction.skill_extractor import extract_skills_from_text

MAX_FILE_SIZE_BYTES_DEFAULT = 5 * 1024 * 1024  # 5 MB fallback if no limit is passed in


class UnsupportedFileError(ValueError):
    pass


class EmptyResumeError(ValueError):
    pass


class CorruptFileError(ValueError):
    pass


@dataclass
class EducationEntry:
    institution: str | None = None
    degree: str | None = None
    field_of_study: str | None = None
    graduation_year: str | None = None
    gpa: str | None = None
    raw_text: str = ""


@dataclass
class ExperienceEntry:
    company: str | None = None
    position: str | None = None
    duration: str | None = None
    raw_text: str = ""


@dataclass
class ProjectEntry:
    name: str | None = None
    raw_text: str = ""


@dataclass
class ResumeProfile:
    raw_text: str
    name: str | None = None
    email: str | None = None
    phone: str | None = None
    linkedin: str | None = None
    github: str | None = None
    portfolio: str | None = None
    summary: str | None = None
    education: list[EducationEntry] = field(default_factory=list)
    experience: list[ExperienceEntry] = field(default_factory=list)
    projects: list[ProjectEntry] = field(default_factory=list)
    skills_raw: list[str] = field(default_factory=list)
    certifications: list[str] = field(default_factory=list)
    achievements: list[str] = field(default_factory=list)
    sections_found: list[str] = field(default_factory=list)


# ---------- raw text extraction ----------

def extract_text_from_pdf(file_bytes: bytes) -> str:
    try:
        import pdfplumber
    except ImportError as e:
        raise RuntimeError("pdfplumber is required for PDF parsing") from e
    try:
        text_parts = []
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            for page in pdf.pages:
                text_parts.append(page.extract_text() or "")
        return "\n".join(text_parts)
    except Exception as e:
        raise CorruptFileError(f"Could not read PDF: {e}") from e


def extract_text_from_docx(file_bytes: bytes) -> str:
    try:
        import docx
    except ImportError as e:
        raise RuntimeError("python-docx is required for DOCX parsing") from e
    try:
        document = docx.Document(io.BytesIO(file_bytes))
        paragraphs = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                paragraphs.extend(cell.text for cell in row.cells)
        return "\n".join(paragraphs)
    except Exception as e:
        raise CorruptFileError(f"Could not read DOCX: {e}") from e


def extract_raw_text(file_bytes: bytes, filename: str, max_file_size_bytes: int = MAX_FILE_SIZE_BYTES_DEFAULT) -> str:
    if len(file_bytes) == 0:
        raise EmptyResumeError("Uploaded file is empty.")
    if len(file_bytes) > max_file_size_bytes:
        raise ValueError(f"File exceeds maximum allowed size of {max_file_size_bytes // (1024 * 1024)}MB.")

    lower = filename.lower()
    if lower.endswith(".pdf"):
        text = extract_text_from_pdf(file_bytes)
    elif lower.endswith(".docx"):
        text = extract_text_from_docx(file_bytes)
    else:
        raise UnsupportedFileError("Only .pdf and .docx files are supported.")

    if not text or not text.strip():
        raise EmptyResumeError("No extractable text found in the resume (it may be a scanned image).")
    return text


# ---------- field extraction ----------

_DEGREE_PATTERN = re.compile(
    r"\b(B\.?Tech|M\.?Tech|B\.?E|M\.?E|B\.?Sc|M\.?Sc|Bachelor(?:'s)?|Master(?:'s)?|Ph\.?D|MBA|BCA|MCA)\b",
    re.IGNORECASE,
)
_YEAR_PATTERN = re.compile(r"\b(19|20)\d{2}\b")
_GPA_PATTERN = re.compile(r"\b(?:CGPA|GPA)\s*[:\-]?\s*(\d\.\d{1,2})\s*(?:/\s*(\d(?:\.\d)?))?", re.IGNORECASE)
_DURATION_PATTERN = re.compile(
    r"\b((?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*\.?\s*\d{4}|\d{4})\s*"
    r"(?:-|–|to)\s*((?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*\.?\s*\d{4}|\d{4}|Present|Current)",
    re.IGNORECASE,
)


def _guess_name(header_text: str) -> str | None:
    """Best-effort: first non-empty line of the header that isn't an email/url/phone."""
    for line in split_lines(header_text)[:5]:
        if "@" in line or "http" in line.lower() or re.search(r"\d{5,}", line):
            continue
        words = line.split()
        if 1 <= len(words) <= 5 and all(w[0].isupper() for w in words if w[0].isalpha()):
            return line.strip()
    return None


def _parse_education_block(block: str) -> list[EducationEntry]:
    entries = []
    for chunk in re.split(r"\n{2,}", block):
        if not chunk.strip():
            continue
        degree_match = _DEGREE_PATTERN.search(chunk)
        year_matches = _YEAR_PATTERN.findall(chunk)
        gpa_match = _GPA_PATTERN.search(chunk)
        lines = split_lines(chunk)
        institution = lines[0] if lines else None
        entries.append(EducationEntry(
            institution=institution,
            degree=degree_match.group(0) if degree_match else None,
            graduation_year=(_YEAR_PATTERN.search(chunk).group(0) if year_matches else None),
            gpa=gpa_match.group(1) if gpa_match else None,
            raw_text=chunk.strip(),
        ))
    return entries


def _parse_experience_block(block: str) -> list[ExperienceEntry]:
    entries = []
    for chunk in re.split(r"\n{2,}", block):
        if not chunk.strip():
            continue
        lines = split_lines(chunk)
        duration_match = _DURATION_PATTERN.search(chunk)
        entries.append(ExperienceEntry(
            position=lines[0] if lines else None,
            company=lines[1] if len(lines) > 1 else None,
            duration=duration_match.group(0) if duration_match else None,
            raw_text=chunk.strip(),
        ))
    return entries


def _parse_projects_block(block: str) -> list[ProjectEntry]:
    entries = []
    for chunk in re.split(r"\n{2,}", block):
        if not chunk.strip():
            continue
        lines = split_lines(chunk)
        entries.append(ProjectEntry(name=lines[0] if lines else None, raw_text=chunk.strip()))
    return entries


def parse_resume(file_bytes: bytes, filename: str, max_file_size_bytes: int = MAX_FILE_SIZE_BYTES_DEFAULT) -> ResumeProfile:
    """Main entry point: bytes + filename -> ResumeProfile."""
    raw_text = extract_raw_text(file_bytes, filename, max_file_size_bytes)
    text = clean_text(raw_text)
    sections = detect_sections(text)
    contact = extract_contact_fields(text)

    profile = ResumeProfile(
        raw_text=text,
        name=_guess_name(sections.get("header", text)),
        email=contact["email"],
        phone=contact["phone"],
        linkedin=contact["linkedin"],
        github=contact["github"],
        portfolio=contact["portfolio"],
        summary=sections.get("summary"),
        sections_found=list(sections.keys()),
    )

    if "education" in sections:
        profile.education = _parse_education_block(sections["education"])
    if "experience" in sections:
        profile.experience = _parse_experience_block(sections["experience"])
    if "projects" in sections:
        profile.projects = _parse_projects_block(sections["projects"])
    if "certifications" in sections:
        profile.certifications = split_lines(sections["certifications"])
    if "achievements" in sections:
        profile.achievements = split_lines(sections["achievements"])

    # Skills: prefer the dedicated skills section, but also scan experience/projects
    # since candidates often only mention tech there, not in a "Skills" list.
    skill_search_text = "\n".join(filter(None, [
        sections.get("skills", ""), sections.get("experience", ""), sections.get("projects", "")
    ]))
    profile.skills_raw = extract_skills_from_text(skill_search_text or text)

    return profile
